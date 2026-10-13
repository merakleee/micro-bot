from dataclasses import dataclass
from pathlib import Path
from tempfile import TemporaryDirectory
from urllib.parse import urlparse
from zipfile import BadZipFile

import httpx
from bs4 import BeautifulSoup
from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from pypdf import PdfReader
from pypdf.errors import PdfReadError


@dataclass(frozen=True)
class ParsedDocument:
    title: str
    text: str


def parse_md(path: Path) -> ParsedDocument:
    try:
        content = path.read_text(encoding="utf-8")
    except (OSError, UnicodeError) as exc:
        raise ValueError(f"Unable to read Markdown file: {path}") from exc

    text = content.replace("\r\n", "\n").replace("\r", "\n").strip()
    if not text:
        raise ValueError(f"Markdown file is empty: {path}")

    return ParsedDocument(title=path.stem, text=text)


def _extract_pdf_pages(path: Path) -> list[str]:
    try:
        reader = PdfReader(path)
        return [(page.extract_text() or "").strip() for page in reader.pages]
    except (OSError, PdfReadError) as exc:
        raise ValueError(f"Unable to read PDF file: {path}") from exc


def parse_pdf(path: Path) -> ParsedDocument:
    try:
        import ocrmypdf
    except ImportError as exc:
        raise ValueError(
            "OCR is required for PDF ingestion. Install OCRmyPDF and Tesseract "
            "with the eng, fra, and ara language data."
        ) from exc

    try:
        with TemporaryDirectory() as temporary_directory:
            ocr_path = Path(temporary_directory) / "ocr.pdf"
            ocrmypdf.ocr(
                str(path),
                str(ocr_path),
                language="eng+fra+ara",
                redo_ocr=True,
                output_type="pdf",
                jobs=1,
                progress_bar=False,
            )
            pages = _extract_pdf_pages(ocr_path)
    except Exception as exc:
        raise ValueError(f"Unable to OCR PDF file: {path}") from exc

    text = "\n\n".join(page for page in pages if page)
    if not text:
        raise ValueError(f"PDF file has no readable text: {path}")

    return ParsedDocument(title=path.stem, text=text)


def parse_docx(path: Path) -> ParsedDocument:
    try:
        document = Document(path)
    except (OSError, BadZipFile, PackageNotFoundError) as exc:
        raise ValueError(f"Unable to read DOCX file: {path}") from exc

    paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs]
    text = "\n".join(paragraph for paragraph in paragraphs if paragraph)
    if not text:
        raise ValueError(f"DOCX file has no readable text: {path}")

    return ParsedDocument(title=path.stem, text=text)


def parse_web_page(url: str) -> ParsedDocument:
    parsed_url = urlparse(url.strip())
    if parsed_url.scheme not in {"http", "https"} or not parsed_url.netloc:
        raise ValueError(f"Invalid web page URL: {url}")

    if parsed_url.hostname in {"microclub.info", "www.microclub.info"}:
        return _parse_rendered_web_page(url.strip())

    try:
        response = httpx.get(url, follow_redirects=True, timeout=10.0)
        response.raise_for_status()
    except httpx.HTTPError as exc:
        raise ValueError(f"Unable to fetch web page: {url}") from exc

    soup = BeautifulSoup(response.text, "html.parser")
    title = soup.title.get_text(" ", strip=True) if soup.title else parsed_url.netloc

    for element in soup.select("script, style, nav, footer, noscript"):
        element.decompose()

    content = soup.body or soup
    lines = [line.strip() for line in content.get_text("\n").splitlines()]
    text = "\n".join(line for line in lines if line)
    if not text:
        raise ValueError(f"Web page has no readable text: {url}")

    return ParsedDocument(title=title, text=text)


def _parse_rendered_web_page(url: str) -> ParsedDocument:
    try:
        from playwright.sync_api import Error as PlaywrightError
        from playwright.sync_api import sync_playwright
    except ImportError as exc:
        raise ValueError(
            "Micro Club pages need Playwright. Install project requirements and "
            "run 'playwright install chromium'."
        ) from exc

    try:
        with sync_playwright() as playwright:
            browser = playwright.chromium.launch(headless=True)
            try:
                page = browser.new_page()
                page.goto(url, wait_until="domcontentloaded", timeout=30_000)
                page.wait_for_selector("#root > *", timeout=20_000)

                page.evaluate(
                    """async () => {
                        const step = Math.max(window.innerHeight, 500);
                        for (let y = 0; y < document.body.scrollHeight; y += step) {
                            window.scrollTo(0, y);
                            await new Promise(resolve => setTimeout(resolve, 120));
                        }
                        window.scrollTo(0, 0);
                    }"""
                )
                page.locator(
                    "script, style, nav, footer, noscript, #Navbar, #Landing > :not(h1)"
                ).evaluate_all(
                    "elements => elements.forEach(element => element.remove())"
                )
                title = page.title().strip() or urlparse(url).netloc
                text = page.locator("body").inner_text().strip()
            finally:
                browser.close()
    except PlaywrightError as exc:
        raise ValueError(f"Unable to render Micro Club page: {url}") from exc

    if not text:
        raise ValueError(f"Micro Club page has no readable text: {url}")

    return ParsedDocument(title=title, text=text)


def parse_document(path: Path) -> ParsedDocument:
    suffix = path.suffix.lower()
    if suffix in {".md", ".markdown"}:
        return parse_md(path)
    if suffix == ".pdf":
        return parse_pdf(path)
    if suffix == ".docx":
        return parse_docx(path)
    raise ValueError(f"Unsupported document format: {suffix or '<none>'}")
