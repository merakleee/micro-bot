import sys
from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document

from app.ingestion import parsers
from app.ingestion.parsers import ParsedDocument


def test_parse_md_returns_title_and_normalized_text(tmp_path: Path):
    path = tmp_path / "membership.md"
    path.write_text("  # Adhésion\n\nMicroClub  ", encoding="utf-8")

    assert parsers.parse_md(path) == ParsedDocument(
        title="membership",
        text="# Adhésion\n\nMicroClub",
    )


def test_parse_md_rejects_empty_file(tmp_path: Path):
    path = tmp_path / "empty.md"
    path.write_text(" \n\t", encoding="utf-8")

    with pytest.raises(ValueError, match="Markdown file is empty"):
        parsers.parse_md(path)


def test_parse_md_rejects_invalid_utf8(tmp_path: Path):
    path = tmp_path / "invalid.md"
    path.write_bytes(b"\xff\xfe")

    with pytest.raises(ValueError, match="Unable to read Markdown"):
        parsers.parse_md(path)


def test_parse_pdf_returns_text_from_pages(tmp_path: Path, monkeypatch):
    path = tmp_path / "events.pdf"
    monkeypatch.setattr(
        parsers,
        "PdfReader",
        lambda _: SimpleNamespace(
            pages=[
                SimpleNamespace(extract_text=lambda: " First page "),
                SimpleNamespace(extract_text=lambda: "Second page"),
            ]
        ),
    )
    monkeypatch.setitem(
        sys.modules,
        "ocrmypdf",
        SimpleNamespace(ocr=lambda _input, output, **_options: Path(output).touch()),
    )

    assert parsers.parse_pdf(path) == ParsedDocument(
        title="events",
        text="First page\n\nSecond page",
    )


def test_parse_pdf_ocr_handles_selectable_and_image_text(tmp_path: Path, monkeypatch):
    path = tmp_path / "mixed.pdf"
    selectable_text = "Text already in the PDF"
    output_text = "Text already in the PDF\nText found in the image"

    def fake_pdf_reader(pdf_path):
        text = output_text if Path(pdf_path).name == "ocr.pdf" else selectable_text
        return SimpleNamespace(pages=[SimpleNamespace(extract_text=lambda: text)])

    def fake_ocr(input_path, output_path, **options):
        assert input_path == str(path)
        assert options["language"] == "eng+fra+ara"
        assert options["redo_ocr"] is True
        Path(output_path).touch()

    monkeypatch.setattr(parsers, "PdfReader", fake_pdf_reader)
    monkeypatch.setitem(
        sys.modules,
        "ocrmypdf",
        SimpleNamespace(ocr=fake_ocr),
    )

    assert parsers.parse_pdf(path) == ParsedDocument(
        title="mixed",
        text=output_text,
    )


def test_parse_docx_returns_non_empty_paragraphs(tmp_path: Path):
    path = tmp_path / "membership.docx"
    document = Document()
    document.add_paragraph("Membership")
    document.add_paragraph("L'adhésion est ouverte aux étudiants.")
    document.save(path)

    assert parsers.parse_docx(path) == ParsedDocument(
        title="membership",
        text="Membership\nL'adhésion est ouverte aux étudiants.",
    )


def test_parse_docx_rejects_empty_document(tmp_path: Path):
    path = tmp_path / "empty.docx"
    Document().save(path)

    with pytest.raises(ValueError, match="DOCX file has no readable text"):
        parsers.parse_docx(path)


@pytest.mark.parametrize(
    ("filename", "parser_name"),
    [
        ("document.md", "parse_md"),
        ("document.markdown", "parse_md"),
        ("document.pdf", "parse_pdf"),
        ("document.docx", "parse_docx"),
    ],
)
def test_parse_document_selects_parser(
    tmp_path: Path, monkeypatch, filename: str, parser_name: str
):
    expected = ParsedDocument(title="document", text="content")
    monkeypatch.setattr(parsers, parser_name, lambda _: expected)

    assert parsers.parse_document(tmp_path / filename) == expected


def test_parse_document_rejects_unsupported_format(tmp_path: Path):
    with pytest.raises(ValueError, match="Unsupported document format"):
        parsers.parse_document(tmp_path / "document.txt")
